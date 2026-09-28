"""Utilidades compartidas por los controladores (exportación a Excel)."""

from __future__ import annotations

import io
from typing import Iterable

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ANCHO_MAXIMO = 60
ANCHO_MINIMO = 12


def dataframe_a_excel(
    df: pd.DataFrame,
    nombre_hoja: str = "Resultados",
    ancho_por_columna: dict[str, int] | None = None,
) -> io.BytesIO:
    """Genera un Excel con encabezado resaltado, filtros y anchos ajustados."""
    buffer = io.BytesIO()
    hoja_nombre = nombre_hoja[:31]

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=hoja_nombre)

        hoja = writer.sheets[hoja_nombre]
        hoja.freeze_panes = "A2"

        relleno = PatternFill("solid", fgColor="0056B3")
        fuente = Font(color="FFFFFF", bold=True)

        for indice, columna in enumerate(df.columns, start=1):
            celda = hoja.cell(row=1, column=indice)
            celda.fill = relleno
            celda.font = fuente
            celda.alignment = Alignment(vertical="center", wrap_text=True)

            ancho = ancho_por_columna.get(str(columna)) if ancho_por_columna else None
            if ancho is None:
                if not df.empty:
                    largo = df[columna].astype(str).str.len().max() or ANCHO_MINIMO
                    ancho = int(max(ANCHO_MINIMO, min(ANCHO_MAXIMO, largo + 2)))
                else:
                    ancho = len(str(columna)) + 2
            hoja.column_dimensions[get_column_letter(indice)].width = ancho

        if not df.empty:
            hoja.auto_filter.ref = hoja.dimensions

    buffer.seek(0)
    return buffer


def nombre_archivo(partes: Iterable[str]) -> str:
    """Construye un nombre de archivo seguro a partir de los filtros usados."""
    limpio = [
        "".join(caracter if caracter.isalnum() or caracter in "-_" else "_" for caracter in str(parte))
        for parte in partes
        if parte
    ]
    return "_".join([parte for parte in limpio if parte]) or "exportacion"
