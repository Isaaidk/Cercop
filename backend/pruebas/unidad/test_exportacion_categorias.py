"""Pruebas de la exportación a Excel por categorías.

Aquí no se comprueba que el archivo se genere —eso ya funcionaba— sino **el reparto en hojas**, que
es lo nuevo y lo que tiene consecuencias si se hace mal:

1. **Una hoja por categoría.** Las ínfimas cuantías y las ofertas no se leen igual ni se trabajan
   igual. Si acabaran mezcladas, quien recibe el archivo tendría que separarlas a mano en cada uso,
   que es exactamente el trabajo que la exportación venía a ahorrar.
2. **El orden de las pestañas es fijo.** Un orden que dependiera de qué categoría apareció primero
   cambiaría las pestañas entre dos archivos con los mismos filtros, y quien lo abre cada mañana
   tendría que buscarlas.
3. **Ninguna fila se pierde.** Es la prueba que más importa: un reparto por categoría que se dejara
   fuera una fuente desconocida produciría un archivo que *parece* completo y no lo está — el peor
   tipo de fallo en algo cuyo propósito es sacar los datos de la plataforma.

Y una que no es de este archivo pero se comprueba aquí porque es la misma decisión vista desde otro
lado: **la categoría forma parte de la clave de caché**. Sin ella, una exportación de ínfimas y otra
de ofertas con los mismos filtros compartirían entrada, y la segunda recibiría el archivo de la
primera.
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from contratacion.aplicacion.casos_uso.exportar_registros import construir_libro
from contratacion.dominio.busqueda import (
    ETIQUETA_OTRAS,
    ETIQUETA_POR_CATEGORIA,
    Categoria,
    Filtros,
    huella_filtros,
)
from contratacion.dominio.errores import DatoInvalido
from contratacion.infraestructura.adaptadores.salida.bd.consultas import _condiciones

GENERADO = datetime(2026, 9, 28, 15, 0, tzinfo=UTC)

# Etiquetas de hoja, sacadas de las constantes y no escritas a mano: si el nombre comercial cambia,
# la prueba tiene que seguir midiendo el reparto y no la ortografía.
HOJA_INFIMAS = ETIQUETA_POR_CATEGORIA[Categoria.INFIMAS]
HOJA_OFERTAS = ETIQUETA_POR_CATEGORIA[Categoria.OFERTAS]
HOJA_CRITERIOS = "Filtros aplicados"


def _fila(fuente: str, codigo: str) -> dict[str, Any]:
    """Una fila con lo mínimo que lee el exportador: su fuente y su código."""
    return {
        "fuente": fuente,
        "codigo": codigo,
        "entidad": "Municipio de prueba",
        "objeto_compra": "Objeto de prueba",
        "fecha_limite_proformas": None,
    }


def _hojas(contenido: bytes) -> list[str]:
    return list(load_workbook(BytesIO(contenido)).sheetnames)


def _libro(filas: list[dict[str, Any]], **cambios: Any) -> Any:
    contenido = construir_libro(filas, filtros=Filtros(**cambios), generado_en=GENERADO)
    return load_workbook(BytesIO(contenido))


# --------------------------------------------------------------------------- #
# Reparto en hojas
# --------------------------------------------------------------------------- #


def test_sin_categoria_el_libro_lleva_una_hoja_por_categoria() -> None:
    """Es el botón «descargar todo»: un solo libro, y cada categoría en su pestaña."""
    contenido = construir_libro(
        [_fila("NCO", "A-1"), _fila("OCDS", "B-1")], filtros=Filtros(), generado_en=GENERADO
    )

    assert _hojas(contenido) == [HOJA_INFIMAS, HOJA_OFERTAS, HOJA_CRITERIOS]


def test_cada_fila_va_a_la_hoja_de_su_categoria() -> None:
    libro = _libro([_fila("NCO", "INF-1"), _fila("OCDS", "OFE-1")])

    infimas = libro[HOJA_INFIMAS]
    ofertas = libro[HOJA_OFERTAS]

    # Fila 1 es la cabecera; los datos empiezan en la 2.
    assert infimas.cell(row=2, column=1).value == "INF-1"
    assert ofertas.cell(row=2, column=1).value == "OFE-1"
    # Y cada hoja lleva **solo** lo suyo.
    assert infimas.max_row == 2
    assert ofertas.max_row == 2


def test_con_categoria_solo_sale_una_hoja() -> None:
    """Es el botón «solo ínfimas» o «solo ofertas».

    La consulta ya vendrá restringida por SQL, así que las filas que llegan son de esa categoría.
    Lo que se comprueba es que la hoja de la **otra** categoría no aparece: un libro con una
    pestaña de ofertas vacía en una descarga de ínfimas confunde a quien lo abre.
    """
    contenido = construir_libro(
        [_fila("NCO", "A-1"), _fila("NCO", "A-2")],
        filtros=Filtros(categoria=Categoria.INFIMAS),
        generado_en=GENERADO,
    )

    assert _hojas(contenido) == [HOJA_INFIMAS, HOJA_CRITERIOS]


def test_una_fila_de_otra_categoria_no_se_esconde_en_la_hoja_pedida() -> None:
    """Si el filtro de la consulta se rompiera, el reparto no debe taparlo.

    Poner esa fila en la hoja de ínfimas la etiquetaría mal —diría que es algo que no es—. Que
    aparezca su propia pestaña es feo y es correcto: deja el fallo a la vista sin perder el dato.
    """
    contenido = construir_libro(
        [_fila("NCO", "A-1"), _fila("OCDS", "B-1")],
        filtros=Filtros(categoria=Categoria.INFIMAS),
        generado_en=GENERADO,
    )

    assert _hojas(contenido) == [HOJA_INFIMAS, ETIQUETA_OTRAS, HOJA_CRITERIOS]


def test_el_orden_de_las_pestanas_no_depende_del_orden_de_las_filas() -> None:
    """Dos archivos con los mismos filtros tienen que tener las pestañas en el mismo sitio."""
    al_derecho = _hojas(
        construir_libro(
            [_fila("NCO", "A-1"), _fila("OCDS", "B-1")], filtros=Filtros(), generado_en=GENERADO
        )
    )
    al_reves = _hojas(
        construir_libro(
            [_fila("OCDS", "B-1"), _fila("NCO", "A-1")], filtros=Filtros(), generado_en=GENERADO
        )
    )

    assert al_derecho == al_reves


def test_una_fuente_desconocida_no_desaparece() -> None:
    """Es la prueba que protege contra el archivo que parece completo y no lo está.

    Un reparto por categoría que solo contemplara NCO y OCDS tiraría en silencio las filas de una
    fuente nueva, y el archivo seguiría abriéndose sin ningún error.
    """
    contenido = construir_libro(
        [_fila("NCO", "A-1"), _fila("NUEVA", "C-1")], filtros=Filtros(), generado_en=GENERADO
    )

    assert _hojas(contenido) == [HOJA_INFIMAS, ETIQUETA_OTRAS, HOJA_CRITERIOS]
    libro = load_workbook(BytesIO(contenido))
    assert libro[ETIQUETA_OTRAS].cell(row=2, column=1).value == "C-1"


def test_una_exportacion_sin_resultados_produce_un_libro_valido() -> None:
    """Excel no sabe abrir un libro sin ninguna hoja, y el botón no puede devolver un error."""
    libro = _libro([], categoria=Categoria.OFERTAS)

    assert HOJA_OFERTAS in libro.sheetnames
    # La cabecera está aunque no haya datos: el archivo sirve para ver con qué filtros se pidió.
    assert libro[HOJA_OFERTAS].cell(row=1, column=1).value == "Código"


def test_la_hoja_de_criterios_dice_la_categoria() -> None:
    libro = _libro([_fila("NCO", "A-1")], categoria=Categoria.INFIMAS)
    criterios = libro[HOJA_CRITERIOS]

    valores = {
        criterios.cell(row=fila, column=1).value: criterios.cell(row=fila, column=2).value
        for fila in range(2, criterios.max_row + 1)
    }

    assert valores.get("Categoría") == HOJA_INFIMAS


# --------------------------------------------------------------------------- #
# La categoría y la caché
# --------------------------------------------------------------------------- #


def test_la_categoria_cambia_la_huella_de_los_filtros() -> None:
    """Sin esto, dos exportaciones distintas compartirían entrada de caché.

    El fallo no daría ningún error: la segunda petición recibiría el archivo de la primera, que es
    correcto para *otra* consulta. Es la misma trampa que ya documentada con `solo_con_plazo`.
    """
    base = Filtros(terminos=("obras",))
    infimas = Filtros(terminos=("obras",), categoria=Categoria.INFIMAS)
    ofertas = Filtros(terminos=("obras",), categoria=Categoria.OFERTAS)

    assert len({huella_filtros(base), huella_filtros(infimas), huella_filtros(ofertas)}) == 3


# --------------------------------------------------------------------------- #
# La categoría en la consulta
# --------------------------------------------------------------------------- #


def _condiciones_de(**cambios: Any) -> tuple[str, dict[str, Any]]:
    condiciones, parametros = _condiciones(Filtros(**cambios))
    return " AND ".join(condiciones), parametros


def test_la_categoria_restringe_la_consulta_a_sus_fuentes() -> None:
    """Sin esta condición, «solo ínfimas» devolvería el histórico entero.

    El reparto en hojas seguiría funcionando y el archivo saldría con dos pestañas: no habría ningún
    error, solo una exportación que ignora lo que se pidió. Es el fallo que esta prueba impide.
    """
    consulta, parametros = _condiciones_de(categoria=Categoria.INFIMAS)

    assert "f.codigo = ANY(:categoria_fuentes)" in consulta
    assert parametros["categoria_fuentes"] == ["NCO"]

    consulta, parametros = _condiciones_de(categoria=Categoria.OFERTAS)
    assert parametros["categoria_fuentes"] == ["OCDS"]


def test_sin_categoria_no_se_añade_ninguna_condicion_por_categoria() -> None:
    consulta, parametros = _condiciones_de(terminos=("obras",))

    assert "categoria_fuentes" not in consulta
    assert "categoria_fuentes" not in parametros


def test_la_categoria_y_el_permiso_se_exigen_los_dos() -> None:
    """Se combinan con «y», y eso es lo correcto.

    Si alguien pide ínfimas y solo tiene concedida una fuente que no las trae, el resultado tiene
    que quedar vacío. Cualquier otra cosa —ampliar por su cuenta, ignorar el permiso— sería servir
    datos que esa persona no puede ver.
    """
    consulta, parametros = _condiciones_de(
        categoria=Categoria.INFIMAS, fuentes_permitidas=("OCDS",)
    )

    assert "f.codigo = ANY(:categoria_fuentes)" in consulta
    assert "f.codigo = ANY(:fuentes_permitidas)" in consulta
    assert parametros["categoria_fuentes"] == ["NCO"]
    assert parametros["fuentes_permitidas"] == ["OCDS"]


# --------------------------------------------------------------------------- #
# Criterios que se contradicen
# --------------------------------------------------------------------------- #


def test_una_categoria_con_una_fuente_ajena_se_rechaza_con_explicacion() -> None:
    """Pedir ínfimas de OCDS no puede dar nada; mejor decirlo que entregar un archivo vacío."""
    with pytest.raises(DatoInvalido) as fallo:
        Filtros(categoria=Categoria.INFIMAS, fuente="OCDS").validado()

    assert "Ínfimas" in str(fallo.value)


def test_una_categoria_con_una_fuente_suya_se_acepta() -> None:
    Filtros(categoria=Categoria.INFIMAS, fuente="nco").validado()


def test_la_fuente_compatible_no_distingue_mayusculas_ni_espacios() -> None:
    Filtros(categoria=Categoria.OFERTAS, fuente=" ocds ").validado()
