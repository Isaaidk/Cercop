"""Pruebas de las columnas elegidas y del análisis de la plantilla.

Dos cosas que se defienden aquí:

1. **Que la selección se cumpla.** Si la empresa elige cuatro columnas, el archivo tiene que salir
   con esas cuatro y ninguna más. Una selección que se ignora en silencio es peor que no ofrecerla:
   alguien decidiría no mandar el nombre del funcionario en un archivo que sí lo lleva.
2. **Que el informe de la plantilla diga la verdad.** Es lo que responde a «¿por qué mis datos no
   caen donde esperaba?», y un informe que se equivoca de hoja manda a arreglar lo que no está roto.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from contratacion.aplicacion.casos_uso.analizar_plantilla import analizar
from contratacion.aplicacion.casos_uso.exportar_registros import (
    COLUMNAS,
    columnas_del_libro,
    construir_libro,
    revisar_columnas,
)
from contratacion.aplicacion.plantillas import HOJA_DE_DATOS
from contratacion.dominio.busqueda import Categoria, Filtros
from contratacion.dominio.errores import DatoInvalido

GENERADO = datetime(2026, 9, 29, 12, 0, tzinfo=UTC)


def _fila(fuente: str, codigo: str, **campos: Any) -> dict[str, Any]:
    fila: dict[str, Any] = {
        "fuente": fuente,
        "codigo": codigo,
        "entidad": "Municipio de prueba",
        "objeto_compra": "Objeto",
        "provincia": "PICHINCHA",
        "canton": "QUITO",
        "fecha_publicacion": datetime(2026, 9, 20, 10, 0, tzinfo=UTC),
        "fecha_limite_proformas": None,
    }
    fila.update(campos)
    return fila


def _plantilla(hojas: dict[str, list[list[Any]]]) -> bytes:
    """Un `.xlsx` con el contenido indicado: una lista de filas por cada hoja."""
    libro = Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _etiquetas(hoja: Any, fila: int = 1) -> list[str]:
    return [
        str(hoja.cell(row=fila, column=columna).value)
        for columna in range(1, hoja.max_column + 1)
        if hoja.cell(row=fila, column=columna).value is not None
    ]


# --------------------------------------------------------------------------- #
# Lo que se puede pedir
# --------------------------------------------------------------------------- #


def test_las_columnas_conocidas_y_su_etiqueta_estan_en_el_catalogo() -> None:
    """Lo que escribe el libro se puede elegir: el catálogo sale de las mismas tuplas."""
    assert revisar_columnas(["codigo", "entidad", "dias_proforma"]) == (
        "codigo",
        "entidad",
        "dias_proforma",
    )


def test_se_quitan_las_repetidas_y_las_vacias() -> None:
    """Dos veces la misma columna es una columna; un hueco no es ninguna."""
    assert revisar_columnas(["codigo", "codigo", " entidad ", ""]) == ("codigo", "entidad")


def test_una_columna_que_no_existe_se_rechaza() -> None:
    """Callarse devolvería un archivo al que le falta justo lo que se pidió, sin ningún aviso."""
    with pytest.raises(DatoInvalido) as fallo:
        revisar_columnas(["codigo", "presupuesto_fantasma"])

    assert "presupuesto_fantasma" in str(fallo.value)


# --------------------------------------------------------------------------- #
# Qué columnas tiene el archivo
# --------------------------------------------------------------------------- #


def test_sin_seleccion_salen_las_conocidas_y_las_que_aparezcan_en_los_datos() -> None:
    filas = [_fila("NCO", "NIC-1", campo_nuevo_de_la_fuente="algo")]
    claves = [clave for clave, _, _ in columnas_del_libro(filas)]
    conocidas = [clave for clave, _, _ in COLUMNAS]

    assert claves[: len(conocidas)] == conocidas
    assert "campo_nuevo_de_la_fuente" in claves


def test_con_seleccion_solo_salen_las_elegidas_y_no_las_extra() -> None:
    """La selección desactiva el añadido automático: quien elige columnas quiere un archivo fijo."""
    filas = [_fila("NCO", "NIC-1", campo_nuevo_de_la_fuente="algo")]
    claves = [clave for clave, _, _ in columnas_del_libro(filas, elegidas=("codigo", "entidad"))]

    assert claves == ["codigo", "entidad"]


def test_el_orden_lo_decide_el_catalogo_y_no_la_peticion() -> None:
    """Elegir «entidad, código» no pone la razón social delante del código."""
    filas = [_fila("NCO", "NIC-1")]
    claves = [clave for clave, _, _ in columnas_del_libro(filas, elegidas=("entidad", "codigo"))]

    assert claves == ["codigo", "entidad"]


def test_el_libro_sin_plantilla_sale_con_las_columnas_elegidas() -> None:
    contenido = construir_libro(
        [_fila("NCO", "NIC-1")],
        filtros=Filtros(fuentes_permitidas=("NCO",)),
        generado_en=GENERADO,
        columnas=("codigo", "entidad"),
    )
    libro = load_workbook(io.BytesIO(contenido))
    hoja = libro[libro.sheetnames[0]]

    assert _etiquetas(hoja) == ["Código", "Razón social"]


def test_en_la_plantilla_lo_no_elegido_queda_vacio() -> None:
    """Su columna sigue ahí —es su diseño— pero sin datos: es exactamente lo que se pidió."""
    plantilla = _plantilla(
        {
            HOJA_DE_DATOS: [
                ["Código", "Razón social", "Objeto de compra"],
            ]
        }
    )
    contenido = construir_libro(
        [_fila("NCO", "NIC-1")],
        filtros=Filtros(fuentes_permitidas=("NCO",)),
        generado_en=GENERADO,
        plantilla=plantilla,
        columnas=("codigo", "objeto_compra"),
    )
    hoja = load_workbook(io.BytesIO(contenido))[HOJA_DE_DATOS]

    assert hoja.cell(row=2, column=1).value == "NIC-1"
    assert hoja.cell(row=2, column=2).value is None
    assert hoja.cell(row=2, column=3).value == "Objeto"


# --------------------------------------------------------------------------- #
# El informe de la plantilla
# --------------------------------------------------------------------------- #


def test_el_informe_dice_a_que_hoja_va_cada_familia() -> None:
    plantilla = _plantilla(
        {
            "ÍNFIMAS": [["Código Necesidad de Contratación", "Tipo de Necesidad"]],
            "OFERTAS": [["Código", "Objeto del Proceso"]],
            "PARTICIPACION": [["Código", "Estado"]],
        }
    )
    informe = analizar(plantilla).como_diccionario()

    assert informe["destinos"] == {"infimas": "ÍNFIMAS", "ofertas": "OFERTAS"}
    assert informe["familias_sin_hoja_propia"] == []


def test_el_informe_avisa_cuando_no_hay_hoja_por_familia() -> None:
    """Es la causa más frecuente de «mis datos no están donde los busco»: una sola hoja de datos."""
    plantilla = _plantilla({HOJA_DE_DATOS: [["Código", "Razón social"]]})
    informe = analizar(plantilla).como_diccionario()

    assert informe["destinos"] == {"infimas": HOJA_DE_DATOS, "ofertas": HOJA_DE_DATOS}
    assert informe["familias_sin_hoja_propia"] == ["infimas", "ofertas"]
    assert informe["hojas"][0]["recibe"] == ["infimas", "ofertas"]


def test_el_informe_localiza_la_fila_de_encabezados_y_las_columnas() -> None:
    plantilla = _plantilla(
        {
            "ÍNFIMAS": [
                ["PUBLICACIONES 2026"],
                ["ÍNFIMA CUANTÍA"],
                ["Tipo de Necesidad", "Código Necesidad de Contratación"],
            ]
        }
    )
    hoja = analizar(plantilla).como_diccionario()["hojas"][0]

    assert hoja["fila_de_encabezado"] == 3
    assert hoja["primera_fila_de_datos"] == 4
    assert [columna["clave"] for columna in hoja["columnas_reconocidas"]] == [
        "tipo_necesidad",
        "codigo",
    ]


def test_el_informe_no_inventa_encabezados_donde_no_los_hay() -> None:
    """Una hoja en blanco no tiene cabecera: se dice, y los datos irían al principio."""
    plantilla = _plantilla({HOJA_DE_DATOS: []})
    hoja = analizar(plantilla).como_diccionario()["hojas"][0]

    assert hoja["fila_de_encabezado"] is None
    assert hoja["columnas_reconocidas"] == []


def test_el_analisis_usa_el_mismo_criterio_que_la_exportacion() -> None:
    """Lo que dice el informe y lo que hace el libro salen del mismo código.

    Se comprueba con la fila donde empiezan los datos: si el informe dijera una y el volcado usara
    otra, el usuario arreglaría la plantilla por un problema que no existe.
    """
    plantilla = _plantilla(
        {
            "ÍNFIMAS": [
                ["PUBLICACIONES 2026"],
                ["Tipo de Necesidad", "Código Necesidad de Contratación"],
            ]
        }
    )
    informe = analizar(plantilla).como_diccionario()["hojas"][0]

    contenido = construir_libro(
        [_fila("NCO", "NIC-1")],
        filtros=Filtros(categoria=Categoria.INFIMAS, fuentes_permitidas=("NCO",)),
        generado_en=GENERADO,
        plantilla=plantilla,
    )
    hoja = load_workbook(io.BytesIO(contenido))["ÍNFIMAS"]
    fila = int(informe["primera_fila_de_datos"])

    assert hoja.cell(row=fila, column=2).value == "NIC-1"
