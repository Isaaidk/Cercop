"""Pruebas de la columna de CPC en la exportación.

La columna responde a la pregunta con la que se trabaja de verdad: «¿qué clasificación tiene esto?».
Se comprueban tres cosas que se rompen por separado: que la columna sale con el CPC resumido, que
los ítems **no** salen volcados como texto de `json` —serían una celda ilegible que además rompe el
ancho de la hoja—, y que la columna se puede elegir desde la pantalla de plantilla como cualquier
otra.
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from contratacion.aplicacion.casos_uso.exportar_registros import (
    CLAVES_EXPORTABLES,
    COLUMNA_CPC,
    construir_libro,
)
from contratacion.dominio.busqueda import Categoria, Filtros

GENERADO = datetime(2026, 9, 29, 10, 0, tzinfo=UTC)
HOJA_INFIMAS = "Ínfimas cuantías"

ITEM = {
    "numero": 1,
    "codigo": "871410032",
    "descripcion_cpc": "LAVADO Y ENGRASADO DE AUTOMOTORES",
    "descripcion": "Lavado, engrasado y pulverizado de la Volqueta 5 kodiak chevrolet",
    "unidad": "Unidad",
    "cantidad": "6.00",
}


def _fila(**cambios: Any) -> dict[str, Any]:
    fila: dict[str, Any] = {
        "fuente": "NCO",
        "codigo": "NIC-0860001560001-2026-00083",
        "entidad": "GADM Rioverde",
        "objeto_compra": "SERVICIO DE LAVADO PULVERIZADO Y ENGRASADO",
        "fecha_limite_proformas": None,
        "items": [],
        "cpc_codigos": [],
    }
    fila.update(cambios)
    return fila


def _hoja(filas: list[dict[str, Any]], **cambios: Any) -> Any:
    contenido = construir_libro(
        filas,
        filtros=Filtros(categoria=Categoria.INFIMAS, **cambios.pop("filtros", {})),
        generado_en=GENERADO,
        **cambios,
    )
    return load_workbook(BytesIO(contenido))[HOJA_INFIMAS]


def _encabezados(hoja: Any) -> list[str]:
    return [celda.value for celda in hoja[1]]


def _valor(hoja: Any, titulo: str, fila: int = 2) -> Any:
    columna = _encabezados(hoja).index(titulo) + 1
    return hoja.cell(row=fila, column=columna).value


def test_la_columna_del_cpc_lleva_el_codigo_y_el_nombre_estandar() -> None:
    hoja = _hoja([_fila(items=[ITEM, ITEM])])

    assert _valor(hoja, COLUMNA_CPC[1]) == "871410032 LAVADO Y ENGRASADO DE AUTOMOTORES"


def test_una_necesidad_con_varios_cpc_los_muestra_todos() -> None:
    otro = {**ITEM, "codigo": "431510128", "descripcion_cpc": "REPUESTOS PARA MOTOR DIESEL"}
    hoja = _hoja([_fila(items=[ITEM, otro])])

    assert _valor(hoja, COLUMNA_CPC[1]) == (
        "871410032 LAVADO Y ENGRASADO DE AUTOMOTORES | 431510128 REPUESTOS PARA MOTOR DIESEL"
    )


def test_sin_ficha_leida_la_celda_queda_vacia() -> None:
    """Vacío es «todavía no se ha leído», que no es lo mismo que «no tiene clasificación»."""
    hoja = _hoja([_fila()])

    assert _valor(hoja, COLUMNA_CPC[1]) is None


def test_los_items_no_salen_volcados_en_una_columna() -> None:
    """Una lista de objetos en una celda es texto de `json` que nadie lee."""
    hoja = _hoja([_fila(items=[ITEM])])
    encabezados = _encabezados(hoja)

    assert "Items" not in encabezados
    assert "Cpc codigos" not in encabezados
    assert not any("descripcion_cpc" in str(valor) for valor in encabezados)


def test_la_descripcion_libre_no_se_cuela_en_la_columna() -> None:
    """La celda muestra la clasificación, no lo que escribió la entidad."""
    hoja = _hoja([_fila(items=[ITEM])])

    assert "Volqueta" not in str(_valor(hoja, COLUMNA_CPC[1]))


def test_la_columna_del_cpc_se_puede_elegir() -> None:
    assert COLUMNA_CPC[0] in CLAVES_EXPORTABLES


def test_elegir_solo_el_cpc_deja_un_archivo_de_una_columna() -> None:
    hoja = _hoja([_fila(items=[ITEM])], columnas=[COLUMNA_CPC[0]])

    assert _encabezados(hoja) == [COLUMNA_CPC[1]]
    assert _valor(hoja, COLUMNA_CPC[1]) == "871410032 LAVADO Y ENGRASADO DE AUTOMOTORES"
