"""Contrasta un listado de Excel del cliente contra lo que hay en nuestra base.

Responde a la pregunta «¿coinciden?» con nombres propios: qué códigos del listado están en el
histórico, cuáles no, y en qué campos discrepa lo que tenemos de lo que dice el listado.

Es un guion de **solo lectura**: no escribe en la base ni toca el Excel. Su razón de ser es que el
listado de la fuente solo publica lo **vigente** y no guarda histórico, así que la única forma de
comprobar una necesidad que ya venció es preguntárselo a nuestra copia; y la única forma de saber si
nuestra copia vale es compararla con el listado que el cliente lleva a mano.

Uso:

    cd backend
    .\\.venv\\Scripts\\python.exe scripts/comparar_listado_excel.py "C:\\ruta\\listado.xlsx"

Opciones (para listados con otra forma):

    --hoja ÍNFIMAS             hoja del libro (por defecto la primera)
    --fila-encabezado 4        fila donde están los títulos
    --columna-codigo 2         columna (1..N) del código de necesidad
    --filtro NCO               fuente de la base con la que se compara (NCO u OCDS)
"""

from __future__ import annotations

import argparse
import asyncio
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))

from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import text  # noqa: E402
from sqlalchemy.ext.asyncio import create_async_engine  # noqa: E402

from contratacion.infraestructura.adaptadores.salida.bd.sesion import (  # noqa: E402
    normalizar_url_bd,
)
from contratacion.infraestructura.config.ajustes import obtener_ajustes  # noqa: E402

# Columna del listado -> campo canónico guardado en `registro.datos`.
EQUIVALENCIAS: tuple[tuple[str, str], ...] = (
    ("Tipo de Necesidad", "tipo_necesidad"),
    ("Fecha de Publicación", "fecha_publicacion"),
    ("Provincia - Cantón", "provincia"),
    ("Descripción del Objeto de compra", "objeto_compra"),
    ("Estado de la Necesidad", "estado"),
    ("Fecha límite para la entrega de proformas", "fecha_limite_proformas"),
    ("Entidad Contratante", "entidad"),
)

# Campos de fecha: se comparan al minuto, no como texto, porque el mismo instante se escribe de
# muchas formas (`2026-09-30 08:00:00`, `2026-09-30T08:00:00`, una celda de fecha de Excel…).
CAMPOS_FECHA = frozenset({"fecha_publicacion", "fecha_limite_proformas"})

# El código se busca en `datos` (y no en una columna) porque es ahí donde vive: el esquema guarda el
# contenido entero de la fuente en `jsonb` y solo saca a columna lo que hace falta para filtrar.
CONSULTA_REGISTROS = """
SELECT r.datos ->> 'codigo' AS codigo,
       r.es_vigente,
       r.datos,
       coalesce(array_length(r.cpc_codigos, 1), 0) AS cpc,
       (SELECT count(*) FROM registro_historial h WHERE h.registro_id = r.id) AS versiones,
       r.primera_vez_visto
FROM registro r
JOIN fuente f ON f.id = r.fuente_id
WHERE f.codigo = :fuente AND r.datos ->> 'codigo' = ANY(:codigos)
"""


def texto(valor: Any) -> str:
    """Normaliza para comparar: sin espacios de más y sin distinguir mayúsculas."""
    return " ".join(str(valor if valor is not None else "").split()).upper()


def fecha_texto(valor: Any) -> str:
    """El instante del listado reducido a `AAAA-MM-DD HH:MM`, que es como se compara."""
    if valor in (None, ""):
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%Y-%m-%d")
    crudo = str(valor).strip().replace("T", " ")
    try:
        return datetime.fromisoformat(crudo).strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return crudo[:16]


def valor_en_base(datos: dict[str, Any], campo: str) -> Any:
    """El valor guardado que corresponde a una columna del listado.

    `Provincia - Cantón` es el único caso con truco: el listado los publica en **una** columna
    (`TUNGURAHUA - CEVALLOS`) y la base los guarda **separados** (`provincia` y `canton`, porque los
    filtros del panel van por provincia y por cantón). Se vuelven a juntar aquí para comparar contra
    la misma forma en que los publica la fuente.
    """
    if campo == "provincia":
        partes = (datos.get("provincia"), datos.get("canton"))
        return " - ".join(str(parte) for parte in partes if parte)
    return datos.get(campo)


def leer_listado(ruta: Path, hoja: str | None, fila: int, columna: int) -> list[dict[str, Any]]:
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        elegida = libro[hoja] if hoja else libro[libro.sheetnames[0]]
        encabezados = [
            str(celda.value).strip() if celda.value is not None else "" for celda in elegida[fila]
        ]
        filas: list[dict[str, Any]] = []
        for cruda in elegida.iter_rows(min_row=fila + 1, values_only=True):
            if len(cruda) < columna or not cruda[columna - 1]:
                continue
            fila_datos = dict(zip(encabezados, cruda, strict=False))
            fila_datos["__codigo__"] = str(cruda[columna - 1]).strip()
            filas.append(fila_datos)
        return filas
    finally:
        libro.close()


async def consultar(codigos: list[str], fuente: str) -> dict[str, dict[str, Any]]:
    motor = create_async_engine(normalizar_url_bd(obtener_ajustes().database_url))
    try:
        async with motor.connect() as conexion:
            resultado = await conexion.execute(
                text(CONSULTA_REGISTROS), {"fuente": fuente, "codigos": codigos}
            )
            filas = resultado.mappings().all()
        return {texto(fila["codigo"]): dict(fila) for fila in filas}
    finally:
        await motor.dispose()


async def principal() -> int:
    analizador = argparse.ArgumentParser(description="Contrasta un listado de Excel con la base.")
    analizador.add_argument("listado", type=Path, help="Ruta del archivo .xlsx")
    analizador.add_argument("--hoja", default=None, help="Hoja del libro (por defecto la primera)")
    analizador.add_argument("--fila-encabezado", type=int, default=4, help="Fila de los títulos")
    analizador.add_argument("--columna-codigo", type=int, default=2, help="Columna del código")
    analizador.add_argument("--filtro", default="NCO", help="Fuente de la base (NCO u OCDS)")
    argumentos = analizador.parse_args()

    if not argumentos.listado.exists():
        print(f"No existe el archivo: {argumentos.listado}")
        return 2

    filas = leer_listado(
        argumentos.listado, argumentos.hoja, argumentos.fila_encabezado, argumentos.columna_codigo
    )
    print(f"Listado: {argumentos.listado.name} · {len(filas)} códigos")
    print(f"Base: fuente {argumentos.filtro}")

    en_base = await consultar([fila["__codigo__"] for fila in filas], argumentos.filtro)

    presentes: list[str] = []
    ausentes: list[str] = []
    con_diferencias: list[str] = []

    for fila in filas:
        codigo = fila["__codigo__"]
        guardado = en_base.get(texto(codigo))
        if guardado is None:
            ausentes.append(codigo)
            continue

        presentes.append(codigo)
        datos = guardado["datos"] or {}
        diferencias: list[tuple[str, str, str]] = []
        for etiqueta, campo in EQUIVALENCIAS:
            valor_listado = fila.get(etiqueta)
            valor_base = valor_en_base(datos, campo)
            if campo in CAMPOS_FECHA:
                iguales = fecha_texto(valor_listado) == fecha_texto(valor_base)
            else:
                iguales = texto(valor_listado) == texto(valor_base)
            if not iguales:
                diferencias.append((etiqueta, str(valor_listado), str(valor_base)))

        estado = "vigente" if guardado["es_vigente"] else "ya no vigente"
        marca = "coincide" if not diferencias else f"{len(diferencias)} diferencias"
        print(
            f"\n{codigo}\n"
            f"   en la base: {estado} · {guardado['cpc']} CPC · "
            f"{guardado['versiones']} versiones · {marca}"
        )
        for etiqueta, valor_listado, valor_base in diferencias:
            print(f"   ≠ {etiqueta}")
            print(f"       listado: {valor_listado}")
            print(f"       base   : {valor_base}")
        if diferencias:
            con_diferencias.append(codigo)

    print("\n" + "=" * 70)
    print(f"En la base: {len(presentes)} de {len(filas)} ({len(con_diferencias)} con diferencias)")
    print(f"Fuera de la base: {len(ausentes)}")
    for codigo in ausentes:
        print(f"   {codigo}")
    if ausentes:
        print(
            "\nLos de arriba no están en el histórico. Si su plazo ya venció, la fuente tampoco"
            "\nlos publica y no se pueden recuperar: son los que se perdieron mientras el worker"
            "\nestuvo parado. Los que sí siguen publicados se capturan en la siguiente vuelta de"
            "\nvigilancia."
        )
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(principal()))
